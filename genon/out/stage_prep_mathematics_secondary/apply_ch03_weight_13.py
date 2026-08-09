#!/usr/bin/env python3
"""Re-weight mathematics IX ch 3 to 13 in the allocation workbook, and pin ch 4.

WHY. Ch 3's summary and mapping were regenerated 2026-08-09 (the old ones carried fabricated
descriptions). The new mapping's effort signals give effort_index 13.0, but the workbook still
carries weight 15 — the pre-regeneration figure. Every other maths·IX chapter has weight ==
effort_index exactly, so ch 3 is the one stale row.

WHAT THIS CHANGES, computed before writing (largest remainder over 210 periods):
    ch 3: 19 -> 17 periods      (the re-weighted chapter)
    ch 4: 14 -> 15 periods      <-- ALREADY AUTHORED, see the pin below
    ch 6: 14 -> 15 periods      (not authored; free to move)
Nothing else moves. Weights are shares of a fixed annual budget, so re-weighting ONE chapter
reprices its siblings — that is the coupling to keep in view, not a side effect of this script.

THE PLACEHOLDERS MOVE TOO, and for a reason. Chapters 9-16 carry weight 10.1875, which is
EXACTLY the mean of the eight real chapters (81.5/8). That is a convention the workbook
encodes, so when ch 3 drops the mean becomes 79.5/8 = 9.9375 and the placeholders follow it.
Verified irrelevant to the outcome: allocating with placeholders held at 10.1875 gives the
identical 16-chapter result, so this preserves the convention without changing any number.
(Social Sciences IX uses a round 12 rather than its mean 11.8889 — the convention is per-combo,
so nothing outside mathematics|IX is touched.)

THE PIN. `canonical_period_pins.json` exists for exactly this: "a chapter whose library is
ALREADY AUTHORED must not have its counts moved by a later change... pinning is the default for
an authored chapter." maths|IX|4 was authored 2026-08-08 at [14, 11, 8] for Rs 106.52 and
certified ALL PASS, so its counts are pinned before the regeneration that would otherwise move
them to the [15, 12, 9] its new period count implies.

RESIDUAL, NOT FIXED HERE (founder call): the pin holds `canonical_periods` but NOT
`recommended_periods`, which becomes 15 for ch 4 while its authored top has 14 units. Nothing
in certification reads `recommended_periods` — check 1 compares files against
`canonical_plan.counts`, which the pin protects — but `canonical_minutes` will read 750 where
the library is 700, and `top_brief_for` would ask for 15 units if ch 4 were ever re-authored
without lifting the pin.

    python3 genon/out/stage_prep_mathematics_secondary/apply_ch03_weight_13.py          # dry run
    python3 genon/out/stage_prep_mathematics_secondary/apply_ch03_weight_13.py --apply

AFTERWARDS, and it is not optional (testing.md 0.3, the runbook pair):
    python3 genon/master_plan.py && python3 genon/variant_plans.py
master_plan.py rebuilds every row from the workbook and WIPES canonical_plan; variant_plans.py
must re-annotate immediately or no row is trustworthy.
"""
import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[3]
WB = ROOT / "data/content/allocation_norms/ncf_chapterwise_period_allocation.xlsx"
PINS = ROOT / "data/content/allocation_norms/canonical_period_pins.json"
BACKUP = ROOT / "backup/allocation_workbook"

SUBJECT, CLS = "Mathematics", "IX"
CH3_OLD, CH3_NEW = 15, 13
PH_OLD, PH_NEW = 10.1875, 9.9375

dry = "--apply" not in sys.argv

# ── the workbook ─────────────────────────────────────────────────────────────
# data_only=False so the Summary sheet's 76 formulas survive the round-trip.
wb = openpyxl.load_workbook(WB, data_only=False)
ws = wb["Chapters"]
edits = []
for r in ws.iter_rows(min_row=2):
    if r[0].value != SUBJECT or r[1].value != CLS:
        continue
    ch, title, cell = r[2].value, str(r[3].value), r[4]
    if ch == 3:
        if cell.value != CH3_OLD:
            raise SystemExit(f"ABORT: ch3 weight is {cell.value!r}, expected {CH3_OLD} — "
                             "the workbook has moved since this script was written.")
        edits.append((cell, ch, CH3_OLD, CH3_NEW))
    elif "Placeholder" in title:
        if cell.value != PH_OLD:
            raise SystemExit(f"ABORT: ch{ch} placeholder weight is {cell.value!r}, "
                             f"expected {PH_OLD}")
        edits.append((cell, ch, PH_OLD, PH_NEW))

if len(edits) != 9:
    raise SystemExit(f"ABORT: expected 9 edits (ch3 + 8 placeholders), found {len(edits)}")

print(f"workbook: {WB.relative_to(ROOT)}")
for cell, ch, old, new in edits:
    print(f"  {cell.coordinate:6} ch{ch:<3} {old} -> {new}")

# ── the pin ──────────────────────────────────────────────────────────────────
pins = json.loads(PINS.read_text(encoding="utf-8"))
key = f"mathematics|{CLS}|4"
already = key in pins["pins"]
print(f"\npin {key}: {'already present' if already else 'ADDING [14, 11, 8]'}")

if dry:
    print("\ndry run — nothing written. Re-run with --apply.")
    raise SystemExit(0)

BACKUP.mkdir(parents=True, exist_ok=True)
ts = datetime.now().strftime("%Y%m%d_%H%M%S")
shutil.copy2(WB, BACKUP / f"ncf_chapterwise_period_allocation_{ts}.xlsx")
shutil.copy2(PINS, BACKUP / f"canonical_period_pins_{ts}.json")
shutil.copy2(ROOT / "data/content/allocation_norms/master_plan.json",
             BACKUP / f"master_plan_{ts}.json")
print(f"\nbacked up workbook + pins + master_plan -> {BACKUP.relative_to(ROOT)}/")

for cell, _ch, _old, new in edits:
    cell.value = new
wb.save(WB)
print(f"wrote {WB.name}")

if not already:
    pins["pins"][key] = {
        "counts": [14, 11, 8],
        "reason": ("S4 pilot chapter, authored 2026-08-08 at [14,11,8] for Rs 106.52 and "
                   "certified ALL PASS. Ch 3's regeneration re-weighted it (effort_index 15 "
                   "-> 13), which reprices every sibling: ch 4 moves 14 -> 15 recommended "
                   "periods and its counts would become [15,12,9], disagreeing with three "
                   "good files on disk. Founder ruling 2026-08-09: pin, do not re-author. "
                   "NOTE the pin holds canonical_periods only — recommended_periods reads 15 "
                   "and canonical_minutes 750 against a 700-minute library. Lift this pin if "
                   "ch 4 is ever re-authored."),
        "date": "2026-08-09",
    }
    PINS.write_text(json.dumps(pins, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"wrote {PINS.name}")

print("\nNEXT (required, testing.md 0.3):")
print("  python3 genon/master_plan.py && python3 genon/variant_plans.py")
